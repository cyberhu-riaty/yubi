"""
重做第2页_三层结构.py —— 按用户给的三层结构图重做汇报 PPT 第 2 页

用户给的参考图（这是他认可的表达方式）：
    横着三摞，每摞一层，层内左侧写"是谁的 + 能不能改"，右侧用一句话说这层干什么：
        我们写的（Python，随便改）
        ├ 控制算法 plugins/     ← 你的主战场
        ├ 界面面板 面板/        ← 监控/参数/录制/辨识/关于
        └ 桥接层               ← 界面↔仿真的数据通道
        官方的 Python 层（实验性，用但盯着它）
        　插件系统 + 19 个现成面板
        官方 C++ 层（16 个 .so，永远不动也拷不走）
        　物理引擎 MuJoCo + 窗口 + Filament 渲染 + ImGui 界面库

本脚本只替换第 2 页，其余页一个像素都不动（用户已在本地改过其它页）。
用法：
    python3 工具/重做第2页_三层结构.py            # 就地改，先自动备份
"""
from __future__ import annotations

import argparse
import os
import shutil

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

这里 = os.path.dirname(os.path.abspath(__file__))
包根 = os.path.dirname(这里)
字体 = "PingFang SC"

色 = {
    "底": RGBColor(0x0D, 0x14, 0x1C),
    "主文": RGBColor(0xDB, 0xE6, 0xF0),
    "次文": RGBColor(0x7B, 0x8F, 0xA3),
    "强调": RGBColor(0x4C, 0xC2, 0xFF),
    "好": RGBColor(0x3D, 0xDC, 0x97),
    "警": RGBColor(0xFF, 0xB4, 0x54),
    # 三层各自的底色：绿（我们）/ 琥珀（官方 Python，实验性）/ 近黑（官方 C++）
    "层1": RGBColor(0x12, 0x28, 0x24),
    "层2": RGBColor(0x24, 0x1E, 0x10),
    "层3": RGBColor(0x07, 0x0A, 0x0E),
    "层1线": RGBColor(0x3D, 0xDC, 0x97),
    "层2线": RGBColor(0xFF, 0xB4, 0x54),
    "层3线": RGBColor(0x5A, 0x6B, 0x7C),
}


def _色(sh, 属性: str):
    try:
        return getattr(sh, 属性).rgb
    except Exception:                                                 # noqa: BLE001
        return None


def 学习配色(页) -> None:
    """沿用现有页的底/主文/强调色，别让新页跟别页有色差。"""
    for sh in 页.shapes:
        if not str(sh.shape_type).startswith("AUTO"):
            continue
        填 = _色(sh, "fill")
        if 填 is not None and sh.width > Inches(10) and sh.height > Inches(6):
            色["底"] = 填
        if 填 is not None and sh.width < Inches(0.2) and sh.height > Inches(6):
            色["强调"] = 填
    for sh in 页.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    try:
                        if r.font.color is not None and r.font.color.rgb is not None:
                            色["主文"] = r.font.color.rgb
                            return
                    except Exception:                                 # noqa: BLE001
                        pass


def 框(页, 左, 上, 宽, 高, 填充, 线, 线宽=1.0, 圆角=0.02):
    sh = 页.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 左, 上, 宽, 高)
    sh.fill.solid()
    sh.fill.fore_color.rgb = 填充
    if 线 is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = 线
        sh.line.width = Pt(线宽)
    sh.shadow.inherit = False
    try:
        sh.adjustments[0] = 圆角
    except Exception:                                                 # noqa: BLE001
        pass
    return sh


def 字(页, 左, 上, 宽, 高, 行们, 对齐=PP_ALIGN.LEFT, 行距=1.08):
    tb = 页.shapes.add_textbox(左, 上, 宽, 高)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = 0
    for i, 条 in enumerate(行们):
        文本, 号, 颜色, 粗 = 条[0], 条[1], 条[2], 条[3]
        等宽 = 条[4] if len(条) > 4 else False
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = 对齐
        p.line_spacing = 行距
        p.space_after = Pt(1.5)
        r = p.add_run()
        r.text = 文本
        r.font.size = Pt(号)
        r.font.color.rgb = 颜色
        r.font.bold = 粗
        r.font.name = "Menlo" if 等宽 else 字体
    return tb


def 造第2页(页, 页宽, 页高) -> None:
    左 = Inches(0.55)
    宽 = Inches(12.20)

    # ---------------- 标题（与其它页同版式）----------------
    字(页, 左, Inches(0.26), 宽, Inches(0.60), [
        ("三层结构：理解这个，就理解了全部", 25, 色["主文"], True),
    ])
    字(页, Inches(0.58), Inches(0.86), 宽, Inches(0.34), [
        ("我们的工作全部在最上面那层；中间那层是官方的扩展口，用但盯着它；最下面那层从来不动。",
         12, 色["次文"], False),
    ])
    线 = 页.shapes.add_shape(MSO_SHAPE.RECTANGLE, 左, Inches(1.20), 宽, Emu(9525))
    线.fill.solid()
    线.fill.fore_color.rgb = RGBColor(0x1E, 0x33, 0x50)
    线.line.fill.background()
    线.shadow.inherit = False

    # 版面说明（改过一版，记下为什么）：
    #   第一版把"右侧注记框"放在 x=9.30，而层内条目一直排到 x=8.25+4.30 —— 两者**压在一起**。
    #   几何检查抓到了这个重叠（它不报错，只是看起来糊成一团）。
    #   现在明确分两栏：左边三摞（x 0.55~8.55），右边一块注记（x 8.75~12.55）——
    #   宽度取自页面右边界 13.33，不再溢出。
    层左, 层宽 = 左, Inches(8.00)
    注左, 注宽 = Inches(8.75), Inches(3.80)
    条目宽 = Inches(5.05)
    注记左 = 层左 + Inches(5.85)
    注记宽 = Inches(2.05)

    层定义 = [
        # (上, 高, 层名, 标语, 线色, 底色, [(图标, 名称, 右侧注记), ...])
        (Inches(1.30), Inches(1.94),
         "我们写的", "Python，随便改", 色["层1线"], 色["层1"],
         [("├─", "控制算法　plugins/ → 算法库/", "你的主战场"),
          ("├─", "界面面板　面板/", "监控/参数/录制/辨识/关于"),
          ("└─", "桥接层　指标桥 / 参数桥 / 辨识桥", "界面 ↔ 仿真的数据通道")]),
        (Inches(3.40), Inches(1.10),
         "官方的 Python 层", "实验性：能用，但要盯着它", 色["层2线"], 色["层2"],
         [("　", "插件系统（消息/注册/通道/句柄）＋ 19 个面板", "我们借它拿窗口与按钮")]),
        (Inches(4.66), Inches(1.26),
         "官方 C++ 层", "16 个 .so，永远不动也拷不走", 色["层3线"], 色["层3"],
         [("　", "物理引擎 MuJoCo ＋ 窗口 ＋ Filament ＋ ImGui", "改一行要重编，所以不改")]),
    ]

    for (上, 高, 层名, 标语, 线色, 底, 条目) in 层定义:
        框(页, 层左, 上, 层宽, 高, 填充=底, 线=线色, 线宽=1.5, 圆角=0.02)
        字(页, 层左 + Inches(0.16), 上 + Inches(0.07), Inches(5.0), Inches(0.30), [
            (f"{层名}（{标语}）", 13.5, 线色, True),
        ])
        for i, (符, 名, 注) in enumerate(条目):
            y = 上 + Inches(0.44) + i * Inches(0.42)
            字(页, 层左 + Inches(0.28), y, 条目宽, Inches(0.38), [
                (f"{符} {名}", 12, 色["主文"], False, True),
            ])
            if 注:
                字(页, 注记左, y, 注记宽, Inches(0.38), [
                    (f"← {注}", 11, 色["次文"], False),
                ])

    # ---------------- 右侧总注记 ----------------
    框(页, 注左, Inches(1.30), 注宽, Inches(1.94),
       填充=RGBColor(0x14, 0x1A, 0x22), 线=RGBColor(0x2A, 0x3A, 0x4A), 线宽=0.75)
    字(页, 注左 + Inches(0.14), Inches(1.38), 注宽 - Inches(0.28), Inches(1.78), [
        ("这张图在说什么", 12, 色["强调"], True),
        ("· 越往下越重：越不该改，也越改不动", 10.5, 色["主文"], False),
        ("· 我们的东西全在最上面那层", 10.5, 色["主文"], False),
        ("· 中间那层是官方给的口子：借它的窗口与面板，但不产生依赖", 10.5, 色["主文"], False),
        ("· 最下面那层是发动机：真物理、真渲染都在这", 10.5, 色["主文"], False),
    ], 行距=1.15)

    # ---------------- 底部：一层一层读下来，结论是什么 ----------------
    字(页, 层左, Inches(6.06), 注宽 + Inches(7.75), Inches(0.34), [
        ("所以我们的策略是：把功能全部长在最上面那层 —— 加算法就是放一个 .py；"
         "中间那层只借不用（脱钩时动一个文件）；最下面那层一行不改。", 11.5, 色["主文"], False),
    ], 行距=1.15)
    字(页, 层左, Inches(6.48), 注宽 + Inches(7.75), Inches(0.60), [
        ("换来：想怎么改就怎么改、加算法不动外壳、出错隔离。", 11, 色["好"], False),
        ("代价：界面受官方框架限制、只有 macOS，而且精力被牵制 ——"
         "最近 22 次提交里 12 次（55%）在跟界面较劲，只有 2 次（9%）在加产品能力。",
         11, 色["次文"], False),
    ], 行距=1.15)


def main() -> int:
    ap = argparse.ArgumentParser(description="按三层结构重做第 2 页")
    ap.add_argument("--文件", default=os.path.join(
        包根, "报告", "07_历史快照", "交付物", "驭臂_工作汇报.pptx"))
    ap.add_argument("--出", default=None)
    ap.add_argument("--页号", type=int, default=2)
    ap.add_argument("--无备份", action="store_true")
    a = ap.parse_args()

    if not os.path.exists(a.文件):
        print(f"✗ 找不到：{a.文件}")
        return 1
    if a.出 is None and not a.无备份:
        备份 = a.文件.replace(".pptx", ".三层图前备份.pptx")
        shutil.copyfile(a.文件, 备份)
        print(f"  ✓ 已备份：{os.path.basename(备份)}")

    prs = Presentation(a.文件)
    页 = prs.slides[a.页号 - 1]
    学习配色(页)
    print(f"  沿用配色：底={色['底']} 主文={色['主文']} 强调={色['强调']}")

    删掉 = 0
    保留背景 = 0
    for sh in list(页.shapes):
        是背景 = (str(sh.shape_type).startswith("AUTO")
                and sh.width > prs.slide_width * 0.9 and sh.height > prs.slide_height * 0.9)
        if 是背景:
            保留背景 += 1
            continue
        sh._element.getparent().remove(sh._element)
        删掉 += 1
    print(f"  清掉旧形状 {删掉} 个，保留背景 {保留背景} 个")

    条 = 页.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.07), prs.slide_height)
    条.fill.solid()
    条.fill.fore_color.rgb = 色["强调"]
    条.line.fill.background()
    条.shadow.inherit = False

    造第2页(页, prs.slide_width, prs.slide_height)

    出 = a.出 or a.文件
    prs.save(出)
    print(f"  ✓ 已写入：{出}")
    print(f"  第 {a.页号} 页现有 {len(prs.slides[a.页号-1].shapes)} 个形状")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
